"""Importa "Asistencia CONCILIA LECTIVOS 25-26.xlsx" a public.concilia_lectivo_usuarios
+ public.concilia_lectivo_horarios.

Una pestana del Excel por centro. Las de lunes a viernes traen columnas
NOMBRE/APELLIDOS/TELEFONO 1/TELEFONO 2/EDAD y luego 10 columnas
Lunes_turno_A..Viernes_turno_B con SI/NO (matricula, no asistencia real).
"SAN PEDRO finde" es un caso especial de fin de semana: Sabado Turno1/Turno2/
Tarde (3 turnos) y Doningo Turno 1/2 (sic, typo del Excel; 2 turnos).

Estrategia: UPSERT (nunca borrado), igual que el resto de import_*.py. Primero
se upsertean los alumnos (clave natural centro+curso+nombre+apellidos) pidiendo
`return=representation` para recuperar sus id, y con esos id se upsertean los
horarios. Reimportar el mismo curso no duplica alumnos ni horarios; si la
matricula de una fila cambia de SI a NO en el Excel, el horario antiguo queda
igual (no se borra) hasta que se revise a mano.

Uso:
  python scripts/import_concilia_lectivo_usuarios.py --dry-run
  python scripts/import_concilia_lectivo_usuarios.py --service-role-key ...
  (o con SUPABASE_SERVICE_ROLE_KEY en el entorno)
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import urllib.error
import urllib.request
from pathlib import Path

import openpyxl

DEFAULT_SOURCE_PATH = Path(
    r"C:\Users\Jr\OneDrive\JR\Tablas maestras\Asistencia CONCILIA LECTIVOS 25-26.xlsx"
)
DEFAULT_SUPABASE_URL = "https://epbtoarkinvgcaewbtvs.supabase.co"
DEFAULT_BATCH_SIZE = 500
CURSO_ESCOLAR = "2025-2026"

# Pestana del Excel -> instalaciones.id (verificado contra la base de datos real).
# El id 221 "CP German Fdez Ramos" (sin tilde) es un duplicado de 204 pendiente
# de retirar; no usarlo aqui.
SHEET_TO_CENTRO_ID = {
    "TRUBIA": 231,               # CP Villar Trubia
    "ERIA": 235,                 # CP La Eria
    "OLLONIEGO": 215,            # CP Narciso Sanchez
    "san claudio": 240,          # CP San Claudio
    "GERMAN FDEZ RAMOS": 204,    # CP German Fernandez Ramos
    "PABLO MIAJA": 213,          # CP Pablo Miaja
    "SAN PEDRO finde": 233,      # CP San Pedro de los Arcos
}

WEEKDAY_SHEETS = {
    "TRUBIA", "ERIA", "OLLONIEGO", "san claudio", "GERMAN FDEZ RAMOS", "PABLO MIAJA",
}

# dia_semana normalizado -> (columna_A, columna_B) tal como aparecen en el Excel,
# resueltas en minusculas al leer el header (el Excel mezcla "Turno"/"turno").
WEEKDAY_COLUMNS = [
    ("lunes", "lunes_turno_a", "lunes_turno_b"),
    ("martes", "martes_turno_a", "martes_turno_b"),
    ("miercoles", "miercoles_turno_a", "miercoles_turno_b"),
    ("jueves", "jueves_turno_a", "jueves_turno_b"),
    ("viernes", "viernes_turno_a", "viernes_turno_b"),
]

# San Pedro finde: (dia_semana, turno, turno_orden, nombre_columna_normalizado)
WEEKEND_COLUMNS = [
    ("sabado", "turno_1", 1, "sabado turno1"),
    ("sabado", "turno_2", 2, "sabado turno2"),
    ("sabado", "tarde", 3, "sabado tarde"),
    ("domingo", "turno_1", 1, "doningo turno 1"),
    ("domingo", "turno_2", 2, "doningo turno 2"),
]


def normalize_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def to_int(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(re.sub(r"[^0-9]", "", str(value)) or 0) or None
    except ValueError:
        return None


def is_si(value) -> bool:
    return str(value or "").strip().lower() == "si"


def read_sheet(ws, centro_id: int) -> tuple[list[dict], list[dict]]:
    rows = list(ws.iter_rows(values_only=True))
    header = [normalize_header(h).replace(" ", "_") for h in rows[0]]
    col_index = {name: i for i, name in enumerate(header) if name}

    usuarios = []
    horarios_by_row = []
    is_weekend = ws.title == "SAN PEDRO finde"

    for row in rows[1:]:
        nombre = row[col_index["nombre"]] if "nombre" in col_index else None
        if nombre is None or not str(nombre).strip():
            continue  # fila de totales (=COUNTIF...) o vacia

        usuario = {
            "centro_id": centro_id,
            "curso_escolar": CURSO_ESCOLAR,
            "nombre": str(nombre).strip(),
            "apellidos": str(row[col_index.get("apellidos", -1)] or "").strip(),
            "telefono_1": str(row[col_index["telefono_1"]] or "").strip() or None
            if "telefono_1" in col_index else None,
            "telefono_2": str(row[col_index["telefono_2"]] or "").strip() or None
            if "telefono_2" in col_index else None,
            "edad": to_int(row[col_index.get("edad", -1)]) if "edad" in col_index else None,
        }

        schedule = []
        if is_weekend:
            for dia_semana, turno, turno_orden, colname in WEEKEND_COLUMNS:
                normalized = colname.replace(" ", "_")
                if normalized in col_index and is_si(row[col_index[normalized]]):
                    schedule.append({"dia_semana": dia_semana, "turno": turno, "turno_orden": turno_orden})
        else:
            for dia_semana, col_a, col_b in WEEKDAY_COLUMNS:
                if col_a in col_index and is_si(row[col_index[col_a]]):
                    schedule.append({"dia_semana": dia_semana, "turno": "A", "turno_orden": 1})
                if col_b in col_index and is_si(row[col_index[col_b]]):
                    schedule.append({"dia_semana": dia_semana, "turno": "B", "turno_orden": 2})

        usuarios.append(usuario)
        horarios_by_row.append(schedule)

    return usuarios, horarios_by_row


def build_records(path: Path) -> tuple[list[dict], list[list[dict]]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    all_usuarios: list[dict] = []
    all_horarios: list[list[dict]] = []

    for sheet_name, centro_id in SHEET_TO_CENTRO_ID.items():
        if sheet_name not in wb.sheetnames:
            print(f"  AVISO: pestana '{sheet_name}' no encontrada en el Excel, se omite.")
            continue
        usuarios, horarios = read_sheet(wb[sheet_name], centro_id)
        print(f"  {sheet_name}: {len(usuarios)} alumnos")
        all_usuarios.extend(usuarios)
        all_horarios.extend(horarios)

    return all_usuarios, all_horarios


def post_batch(supabase_url: str, service_role_key: str, table: str, on_conflict: str,
                batch: list[dict], return_representation: bool) -> list[dict]:
    payload = json.dumps(batch, ensure_ascii=False).encode("utf-8")
    prefer = "resolution=merge-duplicates," + ("return=representation" if return_representation else "return=minimal")
    request = urllib.request.Request(
        f"{supabase_url.rstrip('/')}/rest/v1/{table}?on_conflict={on_conflict}",
        data=payload,
        method="POST",
        headers={
            "apikey": service_role_key,
            "Authorization": f"Bearer {service_role_key}",
            "Content-Type": "application/json",
            "Prefer": prefer,
        },
    )
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            body = response.read()
            return json.loads(body) if return_representation and body else []
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8")
        raise RuntimeError(f"Supabase devolvio {exc.code} en {table}: {detail}") from exc


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    parser = argparse.ArgumentParser(
        description="Importa Asistencia CONCILIA LECTIVOS a concilia_lectivo_usuarios/horarios.")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE_PATH)
    parser.add_argument("--url", default=os.environ.get("SUPABASE_URL", DEFAULT_SUPABASE_URL))
    parser.add_argument("--service-role-key", default=os.environ.get("SUPABASE_SERVICE_ROLE_KEY"))
    parser.add_argument("--batch-size", type=int, default=DEFAULT_BATCH_SIZE)
    parser.add_argument("--dry-run", action="store_true", help="Solo analiza, no escribe en Supabase.")
    args = parser.parse_args()

    print(f"Leyendo {args.source} ...")
    usuarios, horarios_by_row = build_records(args.source)
    total_horarios = sum(len(h) for h in horarios_by_row)
    print(f"\nTotal: {len(usuarios)} alumnos, {total_horarios} filas de horario (curso {CURSO_ESCOLAR}).")

    if args.dry_run:
        print("\n--dry-run: no se escribe nada. Muestra del primer alumno:")
        print(json.dumps(usuarios[0], ensure_ascii=False, indent=2))
        print(json.dumps(horarios_by_row[0], ensure_ascii=False, indent=2))
        return

    if not args.service_role_key:
        raise SystemExit("Falta --service-role-key (o SUPABASE_SERVICE_ROLE_KEY en el entorno).")

    usuario_ids: list[int] = []
    for i in range(0, len(usuarios), args.batch_size):
        batch = usuarios[i:i + args.batch_size]
        returned = post_batch(
            args.url, args.service_role_key, "concilia_lectivo_usuarios",
            "centro_id,curso_escolar,nombre,apellidos", batch, return_representation=True,
        )
        usuario_ids.extend(row["id"] for row in returned)
        print(f"  alumnos {i + len(batch)}/{len(usuarios)}")

    if len(usuario_ids) != len(usuarios):
        raise SystemExit(
            f"Se esperaban {len(usuarios)} id de alumnos y llegaron {len(usuario_ids)}; "
            "revisa a mano antes de importar los horarios."
        )

    horarios = [
        {**slot, "lectivo_usuario_id": usuario_id}
        for usuario_id, slots in zip(usuario_ids, horarios_by_row)
        for slot in slots
    ]
    for i in range(0, len(horarios), args.batch_size):
        batch = horarios[i:i + args.batch_size]
        post_batch(
            args.url, args.service_role_key, "concilia_lectivo_horarios",
            "lectivo_usuario_id,dia_semana,turno", batch, return_representation=False,
        )
        print(f"  horarios {i + len(batch)}/{len(horarios)}")

    print(f"\nListo: {len(usuarios)} alumnos y {len(horarios)} horarios upserted.")


if __name__ == "__main__":
    main()
